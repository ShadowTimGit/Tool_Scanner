import os

from openpyxl.styles import (
    Alignment,
    Border,
    Font,
    PatternFill,
    Side,
)

# Formatting-specific column widths
small_col = 12
med_col = 15
average_col = 18
large_col = 22
file_loc_col = 50

# ==================================================
# Format Header
# ==================================================

def format_header(worksheet):

    header_font = Font(
        bold=True
    )

    for cell in worksheet[1]:
        cell.font = header_font
        cell.alignment = Alignment(
            horizontal="center"
        )


# ==================================================
# Sold Row Formatting
# ==================================================

def format_sold_rows(
    worksheet,
):
    sold_fill = PatternFill(
        fill_type="solid",
        fgColor="C6EFCE",
    )

    sold_column = None

    for cell in worksheet[1]:
        if cell.value == "Sold":
            sold_column = cell.column
            break

    if sold_column is None:
        return

    for row_number in range(
        2,
        worksheet.max_row + 1,
    ):
        sold_value = worksheet.cell(
            row=row_number,
            column=sold_column,
        ).value

        if str(
            sold_value
        ).strip().upper() == "YES":

            # Highlight ONLY the Sold cell,
            # not the entire purchase row.
            worksheet.cell(
                row=row_number,
                column=sold_column,
            ).fill = sold_fill


# ==================================================
# Overview Formatting
# ==================================================

def format_overview_sheet(worksheet):

    column_widths = {
        1: med_col,
        2: average_col,
        3: med_col,
        4: large_col,
        5: med_col,
        6: average_col,
        7: 10,
        8: 30,
        9: 18,
        10: 18,
        11: 18,
    }

    apply_common_formatting(
        worksheet,
        column_widths,
    )

    # ------------------------------------------------
    # Main table formatting
    # ------------------------------------------------

    for row in worksheet.iter_rows(
        min_row=2,
        max_row=worksheet.max_row,
        min_col=1,
        max_col=6,
    ):

        row[2].alignment = Alignment(
            horizontal="left",
            vertical="center",
        )

        row[3].alignment = Alignment(
            horizontal="left",
            vertical="center",
        )

        row[4].alignment = Alignment(
            horizontal="center",
            vertical="center",
        )

        row[5].number_format = "0.00"

        row[5].alignment = Alignment(
            horizontal="right",
            vertical="center",
        )

    # ------------------------------------------------
    # Section styling
    # ------------------------------------------------

    section_fill = PatternFill(
        fill_type="solid",
        fgColor="1F4E78",
    )

    section_font = Font(
        name="Arial",
        size=11,
        bold=True,
        color="FFFFFF",
    )

    label_fill = PatternFill(
        fill_type="solid",
        fgColor="D9EAF7",
    )

    value_fill = PatternFill(
        fill_type="solid",
        fgColor="F2F2F2",
    )

    # ------------------------------------------------
    # Purchase Information
    # ------------------------------------------------

    worksheet.merge_cells("H2:I2")

    for cell_reference in ("H2", "I2"):

        cell = worksheet[cell_reference]

        cell.fill = section_fill
        cell.font = section_font
        cell.alignment = Alignment(
            horizontal="center",
            vertical="center",
        )

    worksheet["H2"] = "PURCHASE INFORMATION"

    worksheet.row_dimensions[2].height = 24

    for cell_reference in (
        "H3",
        "I3",
    ):

        cell = worksheet[cell_reference]

        cell.fill = label_fill
        cell.font = Font(
            name="Arial",
            size=11,
            bold=True,
        )
        cell.alignment = Alignment(
            horizontal="center",
            vertical="center",
            wrap_text=True,
        )

    for cell_reference in (
        "H4",
        "I4",
    ):

        cell = worksheet[cell_reference]

        cell.fill = value_fill
        cell.font = Font(
            name="Arial",
            size=11,
        )
        cell.alignment = Alignment(
            horizontal="center",
            vertical="center",
        )

    # ------------------------------------------------
    # Sales Information
    # ------------------------------------------------

    worksheet.merge_cells("H6:K6")

    for cell_reference in (
        "H6",
        "I6",
        "J6",
        "K6",
    ):

        cell = worksheet[cell_reference]

        cell.fill = section_fill
        cell.font = section_font
        cell.alignment = Alignment(
            horizontal="center",
            vertical="center",
        )

    worksheet["H6"] = "SALES INFORMATION"

    worksheet.row_dimensions[6].height = 24

    for cell_reference in (
        "H7",
        "I7",
        "J7",
        "K7",
    ):

        cell = worksheet[cell_reference]

        cell.fill = label_fill
        cell.font = Font(
            name="Arial",
            size=11,
            bold=True,
        )
        cell.alignment = Alignment(
            horizontal="center",
            vertical="center",
            wrap_text=True,
        )

    for cell_reference in (
        "H8",
        "I8",
        "J8",
        "K8",
    ):

        cell = worksheet[cell_reference]

        cell.fill = value_fill
        cell.font = Font(
            name="Arial",
            size=11,
        )
        cell.alignment = Alignment(
            horizontal="center",
            vertical="center",
        )

    # ------------------------------------------------
    # Storefront Sales Volume
    # ------------------------------------------------

    worksheet.merge_cells("H10:I10")

    for cell_reference in (
        "H10",
        "I10",
    ):

        cell = worksheet[cell_reference]

        cell.fill = section_fill
        cell.font = section_font
        cell.alignment = Alignment(
            horizontal="center",
            vertical="center",
        )

    worksheet["H10"] = "Platform Purchase"

    worksheet.row_dimensions[10].height = 24

    for cell_reference in (
        "H11",
        "H12",
        "H13",
    ):

        cell = worksheet[cell_reference]

        cell.fill = label_fill
        cell.font = Font(
            name="Arial",
            size=11,
            bold=True,
        )
        cell.alignment = Alignment(
            horizontal="center",
            vertical="center",
            wrap_text=True,
        )

    for cell_reference in (
        "I11",
        "I12",
        "I13",
    ):

        cell = worksheet[cell_reference]

        cell.fill = value_fill
        cell.font = Font(
            name="Arial",
            size=11,
        )
        cell.alignment = Alignment(
            horizontal="center",
            vertical="center",
        )

    # ------------------------------------------------
    # Number Formatting
    # ------------------------------------------------

    for cell_reference in (
        "I4",
        "I8",
        "J8",
        "K8",
    ):

        worksheet[cell_reference].number_format = "0.00"

    # ------------------------------------------------
    # Highlight Net Profit
    # ------------------------------------------------

    worksheet["K7"].fill = PatternFill(
        fill_type="solid",
        fgColor="E2F0D9",
    )

    worksheet["K8"].fill = PatternFill(
        fill_type="solid",
        fgColor="E2F0D9",
    )

    worksheet["K7"].font = Font(
        name="Arial",
        size=11,
        bold=True,
    )

    worksheet["K8"].font = Font(
        name="Arial",
        size=11,
        bold=True,
    )


# ==================================================
# Simplified Inventory Formatting
# ==================================================

def format_simplified_inventory_sheet(
    worksheet,
):
    column_widths = {
        1: large_col,
        2: large_col,
        3: med_col,
    }

    apply_common_formatting(
        worksheet,
        column_widths,
    )

    tool_fill = PatternFill(
        fill_type="solid",
        fgColor="D9EAF7",
    )

    tool_font = Font(
        name="Arial",
        size=11,
        bold=True,
    )

    for row_number in range(
        2,
        worksheet.max_row + 1,
    ):
        tool_cell = worksheet.cell(
            row=row_number,
            column=1,
        )

        brand_cell = worksheet.cell(
            row=row_number,
            column=2,
        )

        quantity_cell = worksheet.cell(
            row=row_number,
            column=3,
        )

        # Tool group row
        if (
            tool_cell.value is not None
            and str(tool_cell.value).strip()
        ):
            for column in range(1, 4):
                cell = worksheet.cell(
                    row=row_number,
                    column=column,
                )

                cell.fill = tool_fill
                cell.font = tool_font

        # Quantity
        if isinstance(
            quantity_cell.value,
            (int, float),
        ):
            quantity_cell.number_format = "0"
            quantity_cell.alignment = Alignment(
                horizontal="center",
                vertical="center",
            )

        brand_cell.alignment = Alignment(
            horizontal="left",
            vertical="center",
        )


# ==================================================
# Image Hyperlinks
# ==================================================

def format_image_hyperlinks(
    worksheet,
):
    image_column = None

    for cell in worksheet[1]:
        if cell.value == "File Path to Image":
            image_column = cell.column
            break

    if image_column is None:
        return

    for row_number in range(
        2,
        worksheet.max_row + 1,
    ):
        cell = worksheet.cell(
            row=row_number,
            column=image_column,
        )

        if not cell.value:
            continue

        image_path = str(
            cell.value
        ).strip()

        if not image_path:
            continue

        if not os.path.exists(image_path):
            continue

        cell.hyperlink = image_path
        cell.value = os.path.basename(image_path)
        cell.style = "Hyperlink"


# ==================================================
# JSON Hyperlinks
# ==================================================

def format_json_hyperlinks(
    worksheet,
):
    image_column = None

    for cell in worksheet[1]:
        if cell.value == "File Path to JSON":
            image_column = cell.column
            break

    if image_column is None:
        return

    for row_number in range(
        2,
        worksheet.max_row + 1,
    ):
        cell = worksheet.cell(
            row=row_number,
            column=image_column,
        )

        if not cell.value:
            continue

        json_path = str(
            cell.value
        ).strip()

        if not json_path:
            continue

        if not os.path.exists(json_path):
            continue

        cell.hyperlink = json_path
        cell.value = os.path.basename(json_path)
        cell.style = "Hyperlink"


# ==================================================
# Invoice Total Formatting
# ==================================================

def format_invoice_total_rows(
    worksheet,
):

    total_fill = PatternFill(
        fill_type="solid",
        fgColor="D9EAF7",
    )

    for row_number in range(
        2,
        worksheet.max_row + 1,
    ):

        label_found = False

        for cell in worksheet[row_number]:

            if cell.value == "Invoice Est. Sum":
                label_found = True
                break

        if not label_found:
            continue

        for column in range(
            1,
            worksheet.max_column + 1,
        ):

            cell = worksheet.cell(
                row=row_number,
                column=column,
            )

            cell.font = Font(
                name="Arial",
                size=11,
                bold=True,
            )

            cell.fill = total_fill


# ==================================================
# Common Worksheet Formatting
# ==================================================

def apply_common_formatting(
    worksheet,
    column_widths,
):

    header_fill = PatternFill(
        fill_type="solid",
        fgColor="1F4E78",
    )

    header_font = Font(
        name="Arial",
        size=11,
        bold=True,
        color="FFFFFF",
    )

    body_font = Font(
        name="Arial",
        size=11,
        color="000000",
    )

    highlighted_font = Font(
        name="Arial",
        size=11,
        italic=True,
        color="9C6500",
    )

    thin_border = Border(
        left=Side(style="thin", color="D9D9D9"),
        right=Side(style="thin", color="D9D9D9"),
        top=Side(style="thin", color="D9D9D9"),
        bottom=Side(style="thin", color="D9D9D9"),
    )

    # ------------------------------------------------
    # Column widths
    # ------------------------------------------------

    for column, width in column_widths.items():

        worksheet.column_dimensions[
            worksheet.cell(
                row=1,
                column=column,
            ).column_letter
        ].width = width

    # ------------------------------------------------
    # Header
    # ------------------------------------------------

    for cell in worksheet[1]:
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(
            horizontal="center",
            vertical="center",
        )
        cell.border = thin_border

    worksheet.row_dimensions[1].height = 22

    # ------------------------------------------------
    # Body
    # ------------------------------------------------

    alternate_fill = PatternFill(
        fill_type="solid",
        fgColor="F3F6F9",
    )

    for row_number, row in enumerate(
        worksheet.iter_rows(
            min_row=2,
            max_row=worksheet.max_row,
        ),
        start=2,
    ):
        for cell in row:
            cell.font = body_font
            cell.border = thin_border

            # Alternate normal data rows for readability.
            if row_number % 2 == 0:
                cell.fill = alternate_fill

            if cell.alignment.horizontal:
                horizontal = cell.alignment.horizontal
            else:
                horizontal = None

            cell.alignment = Alignment(
                horizontal=horizontal,
                vertical="center",
            )

    # ------------------------------------------------
    # Remove filters
    # ------------------------------------------------

    worksheet.auto_filter.ref = None

    # ------------------------------------------------
    # Freeze header
    # ------------------------------------------------

    worksheet.freeze_panes = "A2"


# ==================================================
# Purchases Formatting
# ==================================================

def format_purchases_sheet(
    worksheet,
):

    column_widths = {
        1: med_col,
        2: med_col,
        3: large_col,
        4: med_col,
        5: med_col,
        6: large_col,
        7: med_col,
        8: med_col,
        9: small_col,
        10: average_col,
        11: small_col,
        12: average_col,
        13: large_col,
        14: large_col,
        15: large_col,
        16: med_col,
        17: file_loc_col,
        18: file_loc_col,
    }

    apply_common_formatting(
        worksheet,
        column_widths,
    )

    # Size formatting
    size_headers = {
        "Size-1",
        "Size-2",
    }

    for header_cell in worksheet[1]:

        if header_cell.value not in size_headers:
            continue

        column = header_cell.column

        for row_number in range(
            2,
            worksheet.max_row + 1,
        ):

            worksheet.cell(
                row=row_number,
                column=column,
            ).number_format = "@"

    # Center selected columns
    center_headers = {
        "Sold",
        "Specialty Socket",
        "Est. Value",
        "Est. Value Fraction",
    }

    for header_cell in worksheet[1]:

        if header_cell.value not in center_headers:
            continue

        column = header_cell.column

        for row_number in range(
            2,
            worksheet.max_row + 1,
        ):

            worksheet.cell(
                row=row_number,
                column=column,
            ).alignment = Alignment(
                horizontal="center",
                vertical="center",
            )

    # Price/value formatting
    price_headers = {
        #"Invoice Price",
        "Est. Value",
        "Est. Value Fraction",
        "Est. Purchase Price",
    }

    for header_cell in worksheet[1]:

        if header_cell.value not in price_headers:
            continue

        column = header_cell.column

        for row_number in range(
            2,
            worksheet.max_row + 1,
        ):

            cell = worksheet.cell(
                row=row_number,
                column=column,
            )

            cell.number_format = "0.00"


# ==================================================
# Sale Profit Total Formatting
# ==================================================

def format_sale_profit_total_rows(
    worksheet,
):
    total_fill = PatternFill(
        fill_type="solid",
        fgColor="D9EAF7",
    )

    total_font = Font(
        bold=True,
    )

    revenue_labels = {
        "Sales:",
        "Total Sales:",
    }

    profit_labels = {
        "Profit:",
        "Sales Profit:",
        "Total Profit:",
    }

    for row_number in range(
        1,
        worksheet.max_row + 1,
    ):
        row_values = [
            worksheet.cell(
                row=row_number,
                column=column_number,
            ).value
            for column_number in range(
                1,
                worksheet.max_column + 1,
            )
        ]

        if not any(
            value in revenue_labels
            or value in profit_labels
            for value in row_values
        ):
            continue

        # Format the entire summary row
        for column_number in range(
            1,
            worksheet.max_column + 1,
        ):
            cell = worksheet.cell(
                row=row_number,
                column=column_number,
            )

            cell.font = total_font
            cell.fill = total_fill

        # Format the value immediately after each label.
        # This preserves your existing column placement.
        for column_number in range(
            1,
            worksheet.max_column,
        ):
            label = worksheet.cell(
                row=row_number,
                column=column_number,
            ).value

            if (
                label in revenue_labels
                or label in profit_labels
            ):
                value_cell = worksheet.cell(
                    row=row_number,
                    column=column_number + 1,
                )

                value_cell.number_format = "0.00"

                value_cell.alignment = Alignment(
                    horizontal="right",
                    vertical="center",
                )


# ==================================================
# Sales Formatting
# ==================================================

def format_sales_sheet(
    worksheet,
):

    column_widths = {
        1: med_col,
        2: average_col,
        3: large_col,
        4: med_col,
        5: med_col,
        6: med_col,
        7: med_col,
        8: med_col,
        9: small_col,
        10: average_col,
        11: small_col,
        12: average_col,
        13: average_col,
        14: average_col,
        15: average_col,
        16: file_loc_col,
        17: file_loc_col,
    }

    apply_common_formatting(
        worksheet,
        column_widths,
    )

    # Size formatting
    size_headers = {
        "Size-1",
        "Size-2",
    }

    for header_cell in worksheet[1]:

        if header_cell.value not in size_headers:
            continue

        column = header_cell.column

        for row_number in range(
            2,
            worksheet.max_row + 1,
        ):

            worksheet.cell(
                row=row_number,
                column=column,
            ).number_format = "@"

    # Center Inventory Found
    for header_cell in worksheet[1]:

        if header_cell.value != "Inventory Found":
            continue

        column = header_cell.column

        for row_number in range(
            2,
            worksheet.max_row + 1,
        ):

            worksheet.cell(
                row=row_number,
                column=column,
            ).alignment = Alignment(
                horizontal="center",
                vertical="center",
            )

    # Highlight Inventory Found
    found_fill = PatternFill(
        fill_type="solid",
        fgColor="C6EFCE",
    )

    not_found_fill = PatternFill(
        fill_type="solid",
        fgColor="FFC7CE",
    )

    for header_cell in worksheet[1]:

        if header_cell.value != "Inventory Found":
            continue

        column = header_cell.column

        for row_number in range(
            2,
            worksheet.max_row + 1,
        ):

            cell = worksheet.cell(
                row=row_number,
                column=column,
            )

            value = str(
                cell.value
            ).strip().upper()

            if value == "YES":
                cell.fill = found_fill

            elif value == "NO":
                cell.fill = not_found_fill

    # Price formatting
    price_headers = {
        "Invoice Price",
        "Profit",
    }

    for header_cell in worksheet[1]:

        if header_cell.value not in price_headers:
            continue

        column = header_cell.column

        for row_number in range(
            2,
            worksheet.max_row + 1,
        ):

            cell = worksheet.cell(
                row=row_number,
                column=column,
            )

            if isinstance(
                cell.value,
                (int, float),
            ):
                cell.number_format = "0.00"


# ==================================================
# Inventory Formatting
# ==================================================

def format_inventory_sheet(
    worksheet,
):

    column_widths = {
        1: med_col,
        2: med_col,
        3: med_col,
        4: large_col,
        5: med_col,
        6: med_col,
        7: small_col,
        8: med_col,
        9: med_col,
        10: average_col,
        11: large_col,
        12: file_loc_col,
        13: file_loc_col,
    }

    apply_common_formatting(
        worksheet,
        column_widths,
    )

    # Size formatting
    size_headers = {
        "Size-1",
        "Size-2",
    }

    for header_cell in worksheet[1]:

        if header_cell.value not in size_headers:
            continue

        column = header_cell.column

        for row_number in range(
            2,
            worksheet.max_row + 1,
        ):

            worksheet.cell(
                row=row_number,
                column=column,
            ).number_format = "@"

    # Centered Headers
    center_headers = {
        "Specialty Socket",
        "Inventory Found",
        "Est. Value",
    }

    for header_cell in worksheet[1]:

        if header_cell.value not in center_headers:
            continue

        column = header_cell.column

        for row_number in range(
            2,
            worksheet.max_row + 1,
        ):

            worksheet.cell(
                row=row_number,
                column=column,
            ).alignment = Alignment(
                horizontal="center",
                vertical="center",
            )

    # Estimated value formatting
    for header_cell in worksheet[1]:

        if header_cell.value != "Est. Value":
            continue

        column = header_cell.column

        for row_number in range(
            2,
            worksheet.max_row + 1,
        ):

            cell = worksheet.cell(
                row=row_number,
                column=column,
            )

            if isinstance(
                cell.value,
                (int, float),
            ):
                cell.number_format = "0.00"


# ==================================================
# Inventory Total Formatting
# ==================================================

def format_inventory_total_rows(
    worksheet,
):

    total_fill = PatternFill(
        fill_type="solid",
        fgColor="D9EAF7",
    )

    for row_number in range(
        2,
        worksheet.max_row + 1,
    ):

        label = worksheet.cell(
            row=row_number,
            column=10,
        ).value

        if label not in {
            "Est. Profit Sum",
            "Est. Inventory Sum",
        }:
            continue

        for column in range(
            1,
            worksheet.max_column + 1,
        ):

            cell = worksheet.cell(
                row=row_number,
                column=column,
            )

            cell.font = Font(
                name="Arial",
                size=11,
                bold=True,
            )

            cell.fill = total_fill

            if column == 11:
                cell.number_format = "0.00"
                cell.alignment = Alignment(
                    horizontal="center",
                    vertical="center",
                )