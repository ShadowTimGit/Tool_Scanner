import os

from openpyxl import Workbook, load_workbook

from tool_logger.tool_logger_config import (
    PURCHASE_HEADERS,
    SALE_HEADERS,
    INVENTORY_HEADERS,
    SIMPLIFIED_INVENTORY_HEADERS,
    OVERVIEW_HEADERS,
)

from .calculations import (
    calculate_remaining_inventory,
    calculate_purchase_values,
    calculate_sale_profits,
    normalize_invoice,
    purchase_sort_key,
    sort_rows,
)

from .rows import (
    purchase_row,
    sale_row,
    build_purchase_display_rows,
    build_sale_display_rows,
    build_inventory_display_rows,
    build_simplified_inventory_rows,
    build_overview_rows,
    build_purchase_excel_row_lookup,
)

from .workbook_paths import (
    get_workbook_path,
)

from .json_sync import (
    load_invoice_prices,
    get_cell_path_value,
)

from .formatting import (
    format_header,
    format_purchases_sheet,
    format_sales_sheet,
    format_inventory_sheet,
    format_simplified_inventory_sheet,
    format_overview_sheet,
    format_image_hyperlinks,
    format_json_hyperlinks,
    format_sold_rows,
    format_invoice_total_rows,
    format_inventory_total_rows,
    format_sale_profit_total_rows,
)

# ==================================================
# Create Workbook If Needed
# ==================================================

def ensure_workbook():

    workbook_path = get_workbook_path()

    if os.path.exists(workbook_path):
        workbook = load_workbook(workbook_path)

        if "Simplified Inventory" not in workbook.sheetnames:
            simplified_inventory = workbook.create_sheet(
                "Simplified Inventory"
            )

            simplified_inventory.append(
                SIMPLIFIED_INVENTORY_HEADERS
            )

            format_header(
                simplified_inventory
            )

        if "Overview" not in workbook.sheetnames:
            overview = workbook.create_sheet(
                "Overview"
            )

            overview.append(
                OVERVIEW_HEADERS
            )

            format_header(
                overview
            )

        workbook.save(
            workbook_path
        )

        workbook.close()

        return workbook_path

    workbook = Workbook()

    purchases = workbook.active
    purchases.title = "Purchases"

    sales = workbook.create_sheet("Sales")
    inventory = workbook.create_sheet("Remaining Inventory")

    simplified_inventory = workbook.create_sheet(
        "Simplified Inventory"
    )

    overview = workbook.create_sheet(
        "Overview"
    )

    purchases.append(
        PURCHASE_HEADERS
    )

    sales.append(
        SALE_HEADERS
    )

    inventory.append(
        INVENTORY_HEADERS
    )

    simplified_inventory.append(
        SIMPLIFIED_INVENTORY_HEADERS
    )

    overview.append(
        OVERVIEW_HEADERS
    )

    for worksheet in (
        purchases,
        sales,
        inventory,
        simplified_inventory,
        overview,
    ):
        format_header(
            worksheet
        )

    workbook.save(
        workbook_path
    )

    workbook.close()

    return workbook_path


# ==================================================
# Rebuild Entire Workbook
# ==================================================

def rebuild_workbook(
    invoice_prices=None,
):

    workbook_path = ensure_workbook()

    workbook = load_workbook(
        workbook_path
    )

    purchases_sheet = workbook[
        "Purchases"
    ]

    sales_sheet = workbook[
        "Sales"
    ]

    inventory_sheet = workbook[
        "Remaining Inventory"
    ]

    simplified_inventory_sheet = workbook[
        "Simplified Inventory"
    ]

    overview_sheet = workbook[
        "Overview"
    ]

    # ------------------------------------------------
    # Load Existing Data
    # ------------------------------------------------

    purchases = load_sheet_rows(
        purchases_sheet,
        PURCHASE_HEADERS,
    )

    sales = load_sheet_rows(
        sales_sheet,
        SALE_HEADERS,
    )

    if invoice_prices is None:
        invoice_prices = load_invoice_prices(
            purchases_sheet
        )

    # ------------------------------------------------
    # Calculate Inventory
    # ------------------------------------------------

    remaining = calculate_remaining_inventory(
        purchases,
        sales,
    )

    calculate_purchase_values(
        purchases,
        invoice_prices,
    )

    calculate_sale_profits(
        purchases,
        sales,
    )

    # ------------------------------------------------
    # Sort
    # ------------------------------------------------

    purchases.sort(
        key=purchase_sort_key
    )

    remaining = sort_rows(
        remaining
    )

    # Build the exact Excel row lookup after Purchases
    # have been sorted.
    purchase_excel_rows = (
        build_purchase_excel_row_lookup(
            purchases
        )
    )

    # ------------------------------------------------
    # Rebuild Purchases
    # ------------------------------------------------

    rebuild_sheet(
        purchases_sheet,
        PURCHASE_HEADERS,
        build_purchase_display_rows(
            purchases,
            invoice_prices,
        )
    )

    # ------------------------------------------------
    # Rebuild Sales
    # ------------------------------------------------

    rebuild_sheet(
        sales_sheet,
        SALE_HEADERS,
        build_sale_display_rows(
            sales,
            purchase_excel_rows,
        ),
    )

    # ------------------------------------------------
    # Rebuild Inventory
    # ------------------------------------------------

    rebuild_sheet(
        inventory_sheet,
        INVENTORY_HEADERS,
        build_inventory_display_rows(
            remaining
        ),
    )

    # ------------------------------------------------
    # Rebuild Simplified Inventory
    # ------------------------------------------------

    rebuild_sheet(
        simplified_inventory_sheet,
        SIMPLIFIED_INVENTORY_HEADERS,
        build_simplified_inventory_rows(
            remaining
        ),
    )

    # ------------------------------------------------
    # Rebuild Overview
    # ------------------------------------------------

    rebuild_sheet(
        overview_sheet,
        OVERVIEW_HEADERS,
        build_overview_rows(
            purchases,
        ),
    )

    # ------------------------------------------------
    # Formatting
    # ------------------------------------------------

    format_purchases_sheet(
        purchases_sheet
    )

    format_sales_sheet(
        sales_sheet
    )

    format_inventory_sheet(
        inventory_sheet
    )

    format_simplified_inventory_sheet(
        simplified_inventory_sheet
    )

    format_overview_sheet(
        overview_sheet
    )

    format_image_hyperlinks(
        purchases_sheet
    )

    format_image_hyperlinks(
        sales_sheet
    )

    format_json_hyperlinks(
        inventory_sheet
    )

    format_json_hyperlinks(
        purchases_sheet
    )

    format_json_hyperlinks(
        sales_sheet
    )

    format_image_hyperlinks(
        inventory_sheet
    )

    # Apply special row formatting
    format_sold_rows(
        purchases_sheet
    )

    format_invoice_total_rows(
        purchases_sheet
    )

    format_inventory_total_rows(
        inventory_sheet
    )

    format_sale_profit_total_rows(
        sales_sheet
    )

    # ------------------------------------------------
    # Save
    # ------------------------------------------------

    workbook.save(
        workbook_path
    )

    workbook.close()

    return workbook_path

# ==================================================
# Remove a single row from the workbook by JSON/image path
# ==================================================

def remove_row_from_workbook_by_path(
    json_path=None,
    image_path=None,
):
    workbook_path = ensure_workbook()

    if not os.path.exists(workbook_path):
        return False

    workbook = load_workbook(
        workbook_path,
        data_only=False,
    )

    removed = False
    requested_paths = []

    for candidate in (json_path, image_path):
        if candidate is None:
            continue
        value = str(candidate).strip()
        if value and value.lower() != "none":
            requested_paths.append(
                os.path.normcase(
                    os.path.normpath(
                        os.path.abspath(
                            value
                        )
                    )
                )
            )

    for sheet_name in ("Purchases", "Sales"):
        worksheet = workbook[sheet_name]

        header_map = {
            cell.value: cell.column
            for cell in worksheet[1]
            if cell.value is not None
        }

        json_column = header_map.get("File Path to JSON")
        image_column = header_map.get("File Path to Image")

        if json_column is None and image_column is None:
            continue

        for row_number in range(
            worksheet.max_row,
            1,
            -1,
        ):
            match = False

            if json_column is not None:
                json_value = get_cell_path_value(
                    worksheet.cell(
                        row=row_number,
                        column=json_column,
                    )
                )
                if json_value:
                    json_norm = os.path.normcase(
                        os.path.normpath(
                            os.path.abspath(json_value)
                        )
                    )
                    for requested in requested_paths:
                        if json_norm == requested:
                            match = True
                            break
                        if os.path.basename(json_value) == os.path.basename(str(requested).replace('\\','/')):
                            match = True
                            break

            if not match and image_column is not None:
                image_value = get_cell_path_value(
                    worksheet.cell(
                        row=row_number,
                        column=image_column,
                    )
                )
                if image_value:
                    image_norm = os.path.normcase(
                        os.path.normpath(
                            os.path.abspath(image_value)
                        )
                    )
                    for requested in requested_paths:
                        if image_norm == requested:
                            match = True
                            break
                        if os.path.basename(image_value) == os.path.basename(str(requested).replace('\\','/')):
                            match = True
                            break

            if match:
                worksheet.delete_rows(
                    row_number,
                    1,
                )
                removed = True
                break

        if removed:
            break

    if not removed:
        workbook.close()
        return False

    workbook.save(workbook_path)
    workbook.close()

    workbook = load_workbook(
        workbook_path,
        data_only=False,
    )

    if "Purchases" in workbook.sheetnames:
        purchases_sheet = workbook["Purchases"]
        invoice_prices = load_invoice_prices(
            purchases_sheet
        )
    else:
        invoice_prices = {}

    workbook.close()
    rebuild_workbook(invoice_prices)

    return True

# ==================================================
# Rebuild Sheet
# ==================================================

def rebuild_sheet(
    worksheet,
    headers,
    rows,
):

    worksheet.delete_rows(
        1,
        worksheet.max_row,
    )

    worksheet.append(headers)
    format_header(worksheet)

    for row in rows:
        worksheet.append(row)


def save_to_spreadsheet(
    row_data,
):

    workbook_path = ensure_workbook()

    workbook = load_workbook(
        workbook_path
    )

    inventory_type = row_data.get(
        "Inventory Type",
        "Purchases",
    )

    inventory_type = str(
        inventory_type
    ).strip().lower()

    if inventory_type == "sales":

        worksheet = workbook[
            "Sales"
        ]

        worksheet.append(
            sale_row(
                row_data
            )
        )

    else:

        worksheet = workbook[
            "Purchases"
        ]

        worksheet.append(
            purchase_row(
                row_data
            )
        )

    workbook.save(
        workbook_path
    )

    # ------------------------------------------------
    # Preserve all existing invoice prices.
    # ------------------------------------------------

    invoice_prices = load_invoice_prices(
        workbook[
            "Purchases"
        ]
    )

    if inventory_type != "sales":

        invoice = normalize_invoice(
            row_data.get(
                "eBay ID"
            )
        )

        if invoice:

            try:
                invoice_prices[invoice] = float(
                    row_data.get(
                        "Invoice Price",
                        invoice_prices.get(
                            invoice,
                            0.0,
                        ),
                    )
                    or 0.0
                )

            except (
                TypeError,
                ValueError,
            ):
                # If the new row does not contain a valid
                # invoice price, preserve the existing one.
                invoice_prices.setdefault(
                    invoice,
                    0.0,
                )


    # Recalculate Remaining Inventory
    rebuild_workbook(
        invoice_prices
    )

    workbook.close()

    return workbook_path


# ==================================================
# Read Sheet
# ==================================================

def load_sheet_rows(
    worksheet,
    headers,
):
    rows = []

    for row_number in range(
        2,
        worksheet.max_row + 1,
    ):
        values = []

        for column in range(
            1,
            len(headers) + 1,
        ):
            cell = worksheet.cell(
                row=row_number,
                column=column,
            )

            header = headers[column - 1]

            if header in {
                "File Path to JSON",
                "File Path to Image",
            } and cell.hyperlink:
                target = cell.hyperlink.target

                values.append(
                    str(target).strip()
                    if target
                    else cell.value
                )

            else:
                values.append(cell.value)

        if not any(
            value is not None
            and str(value).strip()
            for value in values
        ):
            continue

        # Skip generated summary rows.
        if (
            "Invoice Est. Sum" in values
            or "Sales:" in values
            or "Total Sales:" in values
            or "Profit:" in values
            or "Total Profit" in values
        ):
            continue

        row_data = {}

        for index, header in enumerate(headers):
            row_data[header] = values[index]

        rows.append(row_data)

    return rows


